"""Daily Excel reports: a Summary sheet plus one sheet per account, kept for two weeks."""

import io
import zipfile
from datetime import timedelta

from conftest import MARKET_MORNING, FakeKotak, buy, fake_for, kotak_payload, signup
from openpyxl import load_workbook
from sqlalchemy import select

from app.broker.base import BrokerError, Position
from app.db import session_scope
from app.models import PnlRecord
from app.services import reports, trading

TODAY = MARKET_MORNING.date()


def _account(client, owner, name="Main", ucc="UCC1"):
    return client.post("/api/accounts", json=kotak_payload(name=name, ucc=ucc), headers=owner).json()["id"]


def _me(client, headers):
    return client.get("/api/auth/me", headers=headers).json()["id"]


def _closed(sym, buy_at, sell_at, qty=10):
    """A squared-off intraday position, with P/L as Kotak computes it."""
    return Position("nse_cm", f"{sym}-EQ", "MIS", 0, buy_qty=qty, sell_qty=qty, buy_avg=buy_at, sell_avg=sell_at,
                    pnl=round((sell_at - buy_at) * qty, 2))


def _book(resp):
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(resp.content))


def _wb(rows, day=TODAY, problems=(), data_only=False):
    return load_workbook(io.BytesIO(reports.build_xlsx(rows, day, list(problems))), data_only=data_only)


def _table(ws):
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


def _stocks(ws):
    """(stock, qty, buy, sell, P/L) rows of an account sheet."""
    out = []
    for row in ws.iter_rows(min_row=reports.FIRST_ROW, values_only=True):
        if row[0] is None:
            break
        out.append(tuple(row[:5]))
    return out


def _saved(day=TODAY):
    with session_scope() as db:
        return [(r.account_name, r.symbol, r.buy_qty, r.sell_qty, r.buy_avg, r.sell_avg, r.pnl)
                for r in db.scalars(select(PnlRecord).where(PnlRecord.trade_date == day).order_by(PnlRecord.id))]


# --- the workbook --------------------------------------------------------------------------


def test_summary_then_one_sheet_per_account():
    rows = [
        reports.Row(2, "Wife", "TCS-EQ", -120.5, 5, 5, 3900.0, 3875.9),
        reports.Row(1, "Main", "SBIN-EQ", 88.0, 40, 40, 245.1, 247.3),
        reports.Row(1, "Main", "INFY-EQ", -30.0, 10, 10, 1500.0, 1497.0),
    ]
    wb = _wb(rows)
    assert wb.sheetnames == ["Summary", "Main", "Wife"]

    main = wb["Main"]
    assert main["A1"].value == "Main" and main["A2"].value == "Wed, 14 Oct 2026"
    assert [c.value for c in main[4]][:5] == ["Stock", "Quantity", "Buy price (₹)", "Sell price (₹)", "P/L (₹)"]
    assert _stocks(main) == [("INFY", 10, 1500.0, 1497.0, "=ROUND((D5-C5)*B5,2)"),
                             ("SBIN", 40, 245.1, 247.3, "=ROUND((D6-C6)*B6,2)")]
    assert (main["D8"].value, main["E8"].value) == ("Net profit", "=SUM(E5:E6)")
    assert main["F4"].value is None  # no notes needed, no Note column
    assert main["A5"].font.name == "Arial" and main["E8"].font.bold
    assert "before brokerage" in main["A10"].value

    s = wb["Summary"]
    assert [c.value for c in s[4]] == ["Account", "Stocks", "Net profit (₹)"]
    assert _table(s)[4] == ("Main", "=COUNTA('Main'!A5:A6)", "='Main'!E8")
    assert _table(s)[5] == ("Wife", "=COUNTA('Wife'!A5:A5)", "='Wife'!E7")
    assert s["A5"].hyperlink.location == "'Main'!A1"
    assert (s["A8"].value, s["C8"].value) == ("Total net profit", "=SUM(C5:C6)")
    assert s["C8"].number_format == main["E8"].number_format


def test_numbers_are_stored_with_the_formulas():
    """Previews and Excel's Protected View show stored results without recalculating."""
    rows = [reports.Row(1, "Main", "SBIN-EQ", 88.0, 40, 40, 245.1, 247.3),
            reports.Row(1, "Main", "INFY-EQ", -30.0, 10, 10, 1500.0, 1497.0),
            reports.Row(2, "Wife", "TCS-EQ", 25.5, 10, 0, 2841.2, 0.0)]
    wb = _wb(rows, data_only=True)
    assert [r[4] for r in _stocks(wb["Main"])] == [-30.0, 88.0]
    assert wb["Main"]["E8"].value == 58.0
    assert [r[1:] for r in _table(wb["Summary"])[4:6]] == [(2, 58.0), (1, 25.5)]
    assert wb["Summary"]["C8"].value == 83.5


def test_open_short_and_odd_positions_use_kotaks_pnl_with_a_note():
    rows = [
        reports.Row(1, "Main", "RELIANCE-EQ", 25.5, 10, 0, 2841.2, 0.0, True),     # bought, not sold
        reports.Row(1, "Main", "ONGC-EQ", 0.0, 0, 20, 0.0, 270.9, False),          # shorted, no price
        reports.Row(1, "Main", "SBIN-EQ", 87.0, 40, 40, 245.1, 247.3),             # Kotak's figure differs
    ]
    main = _wb(rows)["Main"]
    assert main["F4"].value == "Note"
    assert _stocks(main) == [("ONGC", 20, None, 270.9, 0.0), ("RELIANCE", 10, 2841.2, None, 25.5),
                             ("SBIN", 40, 245.1, 247.3, 87.0)]
    assert main["F5"].value == "20 shares sold and not bought back today; P/L counts only the shares bought back"
    assert main["F6"].value == "10 shares still open (not sold); P/L uses the last price"
    assert main["F7"].value == "P/L as reported by Kotak Neo"


def test_sheet_names_are_valid_unique_and_safe():
    names = reports.sheet_names(["Main", "main", "Summary", "History", "a/b:c*d?[e]\\f",
                                 "'quoted'", "x" * 40, "x" * 40, "", "O'Brien"])
    assert names == ["Main", "main (2)", "Summary (2)", "History (2)", "a-b-c-d--e--f",
                     "quoted", "x" * 31, "x" * 27 + " (2)", "Account", "O'Brien"]
    assert all(len(n) <= 31 for n in names)

    wb = _wb([reports.Row(1, "=1+1 HUF", "SBIN-EQ", 10.0, 10, 10, 100.0, 101.0),
              reports.Row(2, "O'Brien", "TCS-EQ", 5.0, 1, 1, 10.0, 15.0)])
    tricky = wb["=1+1 HUF"]
    assert tricky["A1"].value == "=1+1 HUF" and tricky["A1"].data_type == "s"  # text, not a formula
    s = wb["Summary"]
    assert s["A5"].data_type == "s"
    assert s["C6"].value == "='O''Brien'!E7"


def test_two_accounts_with_the_same_name_get_their_own_sheets():
    wb = _wb([reports.Row(1, "Main", "SBIN-EQ", 10.0, 10, 10, 100.0, 101.0),
              reports.Row(2, "Main", "TCS-EQ", 5.0, 1, 1, 10.0, 15.0),
              reports.Row(None, "Main", "ITC-EQ", 1.0, 1, 1, 1.0, 2.0)])  # a deleted account
    assert wb.sheetnames == ["Summary", "Main", "Main (2)", "Main (3)"]
    assert [r[0] for r in _table(wb["Summary"])[4:7]] == ["Main", "Main (2)", "Main (3)"]


def test_a_day_without_trades_still_has_a_total():
    s = _wb([], TODAY - timedelta(days=1))["Summary"]
    t = _table(s)
    assert t[4][0] == "No trades on this day" and t[4][2] == 0
    assert t[6][0] == "Total net profit" and t[6][2] == "=SUM(C5:C5)"


# --- reading the broker and saving -------------------------------------------------------------


def test_todays_report_reads_every_account_and_saves_it(client, owner):
    main, wife = _account(client, owner), _account(client, owner, "Wife", "UCC2")
    fake_for(main).positions_map = {
        ("SBIN-EQ", "MIS"): _closed("SBIN", 245.1, 247.3, 40),
        ("INFY-EQ", "MIS"): _closed("INFY", 1500.0, 1497.0),
        ("UNTOUCHED-EQ", "MIS"): Position("nse_cm", "UNTOUCHED-EQ", "MIS", 0),  # no trades: left out
    }
    fake_for(wife).positions_map = {("TCS-EQ", "MIS"): _closed("TCS", 3900.0, 3910.0, 5)}

    wb = _book(client.get(f"/api/reports/{TODAY}.xlsx", headers=owner))
    assert wb.sheetnames == ["Summary", "Main", "Wife"]
    assert _stocks(wb["Main"]) == [("INFY", 10, 1500.0, 1497.0, "=ROUND((D5-C5)*B5,2)"),
                                   ("SBIN", 40, 245.1, 247.3, "=ROUND((D6-C6)*B6,2)")]
    assert _stocks(wb["Wife"])[0][:4] == ("TCS", 5, 3900.0, 3910.0)
    assert ("Main", "SBIN-EQ", 40, 40, 245.1, 247.3, 88.0) in _saved()

    listing = client.get("/api/reports", headers=owner).json()
    assert listing["today"] == TODAY.isoformat() and listing["keep_days"] == 14
    assert listing["oldest"] == (TODAY - timedelta(days=13)).isoformat()
    assert listing["days"] == [{"day": TODAY.isoformat(), "pnl": 108.0, "rows": 3, "accounts": 2}]


def test_open_positions_get_their_pnl_from_the_live_price(client, owner):
    acc = _account(client, owner)
    FakeKotak.prices = {"1594": 105.0}
    fake_for(acc).positions_map = {("INFY-EQ", "MIS"): Position(
        "nse_cm", "INFY-EQ", "MIS", 10, buy_qty=10, buy_avg=100.0, token="1594")}  # Kotak sends no price/P&L

    main = _book(client.get(f"/api/reports/{TODAY}.xlsx", headers=owner))["Main"]
    assert _stocks(main) == [("INFY", 10, 100.0, None, 50.0)]
    assert main["F5"].value == "10 shares still open (not sold); P/L uses the last price"
    live = client.get("/api/accounts/live", headers=owner).json()
    assert live[0]["pnl"] == 50.0  # the dashboard's Today's P&L counts it too


def test_open_position_without_a_price_counts_only_what_was_sold(client, owner):
    acc = _account(client, owner)
    fake = fake_for(acc)

    def no_quotes(_):
        raise BrokerError("quotes down")

    fake.ltp = no_quotes
    fake.positions_map = {("INFY-EQ", "MIS"): Position(
        "nse_cm", "INFY-EQ", "MIS", 4, buy_qty=10, sell_qty=6, buy_avg=100.0, sell_avg=102.0, token="1594")}
    main = _book(client.get(f"/api/reports/{TODAY}.xlsx", headers=owner))["Main"]
    assert _stocks(main) == [("INFY", 10, 100.0, 102.0, 12.0)]  # 6 sold × ₹2
    assert main["F5"].value.endswith("P/L counts only the shares already sold")


def test_earlier_days_come_from_what_was_saved(client, owner, clock):
    acc = _account(client, owner)
    fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 800.0, 850.0)}
    assert client.get(f"/api/reports/{TODAY}.xlsx", headers=owner).status_code == 200

    clock.set(MARKET_MORNING + timedelta(days=1))
    fake_for(acc).positions_map = {}  # the broker forgets yesterday overnight
    wb = _book(client.get(f"/api/reports/{TODAY}.xlsx", headers=owner))
    assert _stocks(wb["Main"])[0][:4] == ("SBIN", 10, 800.0, 850.0)
    assert wb["Summary"]["A7"].value == "Total net profit"
    assert client.get(f"/api/reports/{TODAY}.xlsx", headers=owner).headers["x-report-warnings"] == "0"


def test_a_broker_error_keeps_the_last_saved_numbers(client, owner):
    acc = _account(client, owner)
    fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 100.0, 107.5)}
    client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)

    def down():
        raise BrokerError("Kotak is down")

    fake_for(acc).positions = down
    r = client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)
    wb = _book(r)
    assert _stocks(wb["Main"])[0][:4] == ("SBIN", 10, 100.0, 107.5)
    assert r.headers["x-report-warnings"] == "1"  # the app warns after the download
    problem = _table(wb["Summary"])[5][0]  # right under the accounts, in red
    assert problem.startswith("Not included: Main: couldn't read positions")
    assert wb["Summary"]["A6"].font.color.rgb.endswith("B42318")


def test_saved_pnl_survives_kotak_clearing_the_day(client, owner):
    acc = _account(client, owner)
    fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 100.0, 107.5)}
    client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)
    fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 100.0, 108.0)}
    client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)  # later numbers replace earlier ones

    fake_for(acc).positions_map = {}  # evening: Kotak no longer lists the day's positions
    wb = _book(client.get(f"/api/reports/{TODAY}.xlsx", headers=owner))
    assert _stocks(wb["Main"])[0][:4] == ("SBIN", 10, 100.0, 108.0)
    assert client.get("/api/reports", headers=owner).json()["days"][0]["rows"] == 1


def test_square_off_saves_the_days_trades(client, owner, custom_sizing):
    acc = _account(client, owner)
    client.post("/api/basket", json={"symbol": "SBIN", "trading_symbol": "SBIN-EQ", "token": "3045",
                                     "sizing_value": 5}, headers=owner)
    buy(client, owner)
    pos = fake_for(acc).positions_map[("SBIN-EQ", "MIS")]
    pos.buy_qty, pos.buy_avg, pos.pnl = 5, 100.0, 42.0
    trading.run_exit(_me(client, owner), "manual")
    assert [(r[0], r[1], r[2], r[6]) for r in _saved()] == [("Main", "SBIN-EQ", 5, 42.0)]


def test_evening_job_saves_only_accounts_tradelab_traded(client, owner, custom_sizing):
    traded, idle = _account(client, owner), _account(client, owner, "Idle", "UCC2")
    client.post("/api/basket", json={"symbol": "SBIN", "trading_symbol": "SBIN-EQ", "token": "3045",
                                     "sizing_value": 5}, headers=owner)
    client.patch(f"/api/accounts/{idle}", json={"automation_enabled": False}, headers=owner)
    buy(client, owner)
    pos = fake_for(traded).positions_map[("SBIN-EQ", "MIS")]
    pos.buy_qty, pos.pnl = 5, 10.0
    fake_for(idle).positions_map = {("TCS-EQ", "MIS"): _closed("TCS", 1.0, 2.0)}

    reports.capture_traded_accounts()
    assert [(r[0], r[1]) for r in _saved()] == [("Main", "SBIN-EQ")]


def test_reports_are_private(client, owner):
    acc = _account(client, owner)
    fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 100.0, 150.0)}
    client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)

    other = {"Authorization": f"Bearer {signup(client, 'friend@example.com').json()['token']}"}
    assert client.get("/api/reports", headers=other).json()["days"] == []
    wb = _book(client.get(f"/api/reports/{TODAY}.xlsx", headers=other))
    assert wb.sheetnames == ["Summary"] and wb["Summary"]["A5"].value == "No trades on this day"
    assert client.get(f"/api/reports/{TODAY}.xlsx").status_code == 401
    assert client.get("/api/reports/all.zip").status_code == 401


def test_bad_or_future_dates_are_rejected(client, owner):
    assert client.get("/api/reports/yesterday.xlsx", headers=owner).status_code == 422
    r = client.get(f"/api/reports/{TODAY + timedelta(days=1)}.xlsx", headers=owner)
    assert r.status_code == 404


# --- two weeks, then gone ------------------------------------------------------------------------


def _old_record(user_id, days_ago, symbol="OLD-EQ"):
    with session_scope() as db:
        db.add(PnlRecord(user_id=user_id, trade_date=TODAY - timedelta(days=days_ago), account_id=None,
                         account_name="Main", symbol=symbol, pnl=1.0, buy_qty=1, sell_qty=1,
                         buy_avg=1.0, sell_avg=2.0))


def test_reports_older_than_two_weeks_are_deleted_everywhere(client, owner):
    user = _me(client, owner)
    _old_record(user, 13)  # still within the two weeks
    _old_record(user, 14)  # 14 days old: deleted
    _old_record(user, 30)

    days = [d["day"] for d in client.get("/api/reports", headers=owner).json()["days"]]
    assert days == [(TODAY - timedelta(days=13)).isoformat()]  # never listed, even before the purge runs
    r = client.get(f"/api/reports/{TODAY - timedelta(days=14)}.xlsx", headers=owner)
    assert r.status_code == 404 and "deleted automatically" in r.json()["detail"]
    assert client.get(f"/api/reports/{TODAY - timedelta(days=13)}.xlsx", headers=owner).status_code == 200

    assert reports.purge_old() == 2
    with session_scope() as db:
        assert [r.trade_date for r in db.scalars(select(PnlRecord))] == [TODAY - timedelta(days=13)]


def test_saving_today_also_clears_out_old_reports(client, owner):
    _account(client, owner)
    _old_record(_me(client, owner), 20)
    client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)
    with session_scope() as db:
        assert db.scalars(select(PnlRecord)).all() == []


def test_the_oldest_day_drops_off_as_days_pass(client, owner, clock):
    _old_record(_me(client, owner), 13)
    assert len(client.get("/api/reports", headers=owner).json()["days"]) == 1
    clock.set(MARKET_MORNING + timedelta(days=1))
    assert client.get("/api/reports", headers=owner).json()["days"] == []
    assert reports.purge_old() == 1


# --- download all ---------------------------------------------------------------------------------


def test_download_all_zips_every_kept_day(client, owner):
    acc = _account(client, owner)
    user = _me(client, owner)
    _old_record(user, 2, "ITC-EQ")
    _old_record(user, 5, "TCS-EQ")
    _old_record(user, 15, "GONE-EQ")  # past the two weeks: not included
    fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 100.0, 101.0)}

    r = client.get("/api/reports/all.zip", headers=owner)
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    first, last = TODAY - timedelta(days=5), TODAY
    assert r.headers["content-disposition"] == f'attachment; filename="TradeLab-PnL-{first}-to-{last}.zip"'
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    assert zf.namelist() == [f"TradeLab-PnL-{TODAY - timedelta(days=5)}.xlsx",
                             f"TradeLab-PnL-{TODAY - timedelta(days=2)}.xlsx",
                             f"TradeLab-PnL-{TODAY}.xlsx"]
    today = load_workbook(io.BytesIO(zf.read(f"TradeLab-PnL-{TODAY}.xlsx")))
    assert _stocks(today["Main"])[0][:4] == ("SBIN", 10, 100.0, 101.0)  # today's is read live
    older = load_workbook(io.BytesIO(zf.read(f"TradeLab-PnL-{TODAY - timedelta(days=5)}.xlsx")))
    assert _stocks(older["Main"])[0][0] == "TCS"


def test_download_all_with_nothing_saved_still_gives_today(client, owner):
    r = client.get("/api/reports/all.zip", headers=owner)
    assert r.status_code == 200
    assert zipfile.ZipFile(io.BytesIO(r.content)).namelist() == [f"TradeLab-PnL-{TODAY}.xlsx"]
    assert r.headers["content-disposition"].endswith(f'filename="TradeLab-PnL-{TODAY}.zip"')


# --- review findings -----------------------------------------------------------------------------


def test_same_stock_under_two_products_is_one_row_and_breaks_nothing(client, owner):
    """Kotak lists one row per product: intraday SBIN plus delivery SBIN in the same account."""
    main, wife = _account(client, owner), _account(client, owner, "Wife", "UCC2")
    mis = _closed("SBIN", 100.0, 110.0, 10)                               # +100
    cnc = Position("nse_cm", "SBIN-EQ", "CNC", -5, sell_qty=5, sell_avg=112.0, pnl=0.0)  # sold from holdings
    fake_for(main).positions_map = {("SBIN-EQ", "MIS"): mis}
    fake_for(wife).positions_map = {("TCS-EQ", "MIS"): _closed("TCS", 10.0, 11.0)}
    assert client.get(f"/api/reports/{TODAY}.xlsx", headers=owner).status_code == 200  # saved once already

    fake_for(main).positions_map[("SBIN-EQ", "CNC")] = cnc  # shows up later in the day
    r = client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)
    wb = _book(r)
    assert r.headers["x-report-warnings"] == "0"
    (sbin,) = _stocks(wb["Main"])
    assert sbin[:2] == ("SBIN", 15) and sbin[2] == 100.0 and abs(sbin[3] - (1100 + 560) / 15) < 1e-9
    assert sbin[4] == 100.0
    assert wb["Main"]["F5"].value.startswith("5 shares sold and not bought back today")
    assert [(x[0], x[1], x[2], x[3]) for x in _saved()] == [("Main", "SBIN-EQ", 10, 15), ("Wife", "TCS-EQ", 10, 10)]
    assert client.get("/api/reports/all.zip", headers=owner).status_code == 200


def test_merging_products_adds_up_quantities_amounts_and_pnl():
    parts = [
        Position("nse_cm", "SBIN-EQ", "MIS", 0, buy_qty=10, sell_qty=10, buy_avg=100.0, sell_avg=101.0, pnl=10.0),
        Position("nse_cm", "SBIN-EQ", "CNC", 5, buy_qty=5, buy_avg=106.0, pnl=None),  # open, no price
    ]
    (row,) = reports.rows_from_positions(1, "Main", parts)
    assert (row.buy_qty, row.sell_qty, row.pnl, row.pnl_complete) == (15, 10, 10.0, False)
    assert abs(row.buy_avg - (1000 + 530) / 15) < 1e-9 and row.sell_avg == 101.0


def test_fallback_pnl_uses_the_contract_multiplier():
    fut = Position("cde_fo", "USDINR26OCTFUT", "NRML", 1, buy_qty=2, sell_qty=1, buy_avg=83.1, sell_avg=83.2,
                   multiplier=1000)
    (row,) = reports.rows_from_positions(1, "Main", [fut])
    assert row.pnl == 100.0 and row.pnl_complete is False


def test_two_saves_at_the_same_moment_both_succeed(client, owner):
    import threading

    accounts = [_account(client, owner, f"A{i}", f"UCC{i}") for i in range(4)]
    for acc in accounts:
        fake_for(acc).positions_map = {(f"S{j}-EQ", "MIS"): _closed(f"S{j}", 10.0, 11.0) for j in range(10)}
    user, results = _me(client, owner), []
    ts = [threading.Thread(target=lambda: results.append(reports.capture(user)[1])) for _ in range(3)]
    for t in ts:
        t.start()
    for t in ts:
        t.join(10)
    assert results == [[], [], []]
    assert len(_saved()) == 40


def test_tabs_sort_numbered_accounts_naturally():
    rows = [reports.Row(i, f"Account {i}", "SBIN-EQ", 1.0, 1, 1, 1.0, 2.0) for i in (10, 2, 1, 15, 3)]
    assert _wb(rows).sheetnames == ["Summary", "Account 1", "Account 2", "Account 3", "Account 10", "Account 15"]


def test_tab_names_fit_excels_limit_with_emoji():
    (name,) = reports.sheet_names(["🚀 Rocket 🚀 account with emoji names"])
    assert len(name.encode("utf-16-le")) // 2 <= 31
    assert reports.sheet_names(["🚀" * 20, "🚀" * 20]) == ["🚀" * 15, "🚀" * 13 + " (2)"]


def test_days_saved_by_the_old_version_say_so():
    main = _wb([reports.Row(1, "Main", "SBIN-EQ", 88.0)])["Main"]
    assert _stocks(main) == [("SBIN", None, None, None, 88.0)]
    assert main["F5"].value == "Saved before quantities and prices were recorded"


def test_account_count_tells_apart_accounts_with_the_same_name(client, owner):
    a, b = _account(client, owner, "Ravi", "UCC1"), _account(client, owner, "Ravi", "UCC2")
    for acc in (a, b):
        fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 10.0, 11.0)}
    client.get(f"/api/reports/{TODAY}.xlsx", headers=owner)
    assert client.get("/api/reports", headers=owner).json()["days"][0]["accounts"] == 2


def test_weekends_and_early_mornings_never_read_the_broker(client, owner, clock):
    acc = _account(client, owner)
    _old_record(_me(client, owner), 3, "ITC-EQ")
    fake_for(acc).positions_map = {("SBIN-EQ", "MIS"): _closed("SBIN", 100.0, 101.0)}  # last session's
    calls = []
    real = fake_for(acc).positions
    fake_for(acc).positions = lambda: calls.append(1) or real()

    for when in (MARKET_MORNING.replace(hour=8, minute=30),          # trading day, before 09:00
                 MARKET_MORNING.replace(day=17, hour=11)):            # Saturday
        clock.set(when)
        listing = client.get("/api/reports", headers=owner).json()
        assert listing["today_live"] is False
        wb = _book(client.get(f"/api/reports/{when.date()}.xlsx", headers=owner))
        assert wb.sheetnames == ["Summary"]
        names = zipfile.ZipFile(io.BytesIO(client.get("/api/reports/all.zip", headers=owner).content)).namelist()
        assert f"TradeLab-PnL-{when.date()}.xlsx" not in names
    assert calls == []
    assert _saved(MARKET_MORNING.date().replace(day=17)) == []


def test_download_all_with_nothing_to_download(client, owner, clock):
    clock.set(MARKET_MORNING.replace(day=17, hour=11))  # Saturday, nothing saved
    r = client.get("/api/reports/all.zip", headers=owner)
    assert r.status_code == 404 and "no reports" in r.json()["detail"]
